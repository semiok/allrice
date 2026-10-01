#!/usr/bin/env python3
"""Bounded macOS experiment, not a Bridge execution backend or Office release.

Use an explicitly prepared Python environment with openpyxl 3.1.5. The upstream
guide and checker remain unchanged. Evidence stays outside the source tree.
"""

import argparse
import hashlib
import json
import os
from pathlib import Path
import platform
import signal
import subprocess
import sys
import time


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--python', required=True, type=Path)
    parser.add_argument('--evidence', required=True, type=Path)
    args = parser.parse_args()
    if sys.platform != 'darwin':
        parser.error('This feasibility experiment requires macOS; it does not certify other platforms.')
    root = Path(__file__).resolve().parents[3]
    python = args.python.absolute()
    directory = args.evidence.resolve()
    directory.mkdir(parents=True, exist_ok=False)
    work = directory / '中文任务'
    inputs, outputs, temporary = [work / name for name in ('input', 'output', 'tmp')]
    for path in (inputs, outputs, temporary):
        path.mkdir(parents=True)
    source = inputs / '原始对账.xlsx'
    target = outputs / '验证结果.xlsx'
    sentinel = directory / 'outside-output.txt'
    sentinel.write_text('preserved', encoding='utf-8')
    (outputs / 'escape-link').symlink_to(sentinel)
    checker = root / 'skills/office/scripts/check_office.py'
    guide = root / 'skills/office/references/xlsx.md'
    env = {'PATH': '/usr/bin:/bin', 'LANG': 'en_US.UTF-8',
           'TMPDIR': str(temporary), 'PYTHONDONTWRITEBYTECODE': '1'}
    setup = subprocess.run([str(python), '-c', '''
import json, platform, sys, sysconfig
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font
w=Workbook(); s=w.active; s.title='对账明细'
s.append(['编号', '金额', '合计']); s.append(['00123', 30, '=B2*2'])
s['A1'].font=Font(bold=True)
w.save(sys.argv[1])
print(json.dumps({'python':sys.version.split()[0], 'prefix':sys.prefix,
 'basePrefix':sys.base_prefix, 'purelib':sysconfig.get_path('purelib'),
 'openpyxl':openpyxl.__version__, 'machine':platform.machine()}))
''', str(source)], env=env, check=True, capture_output=True, text=True, timeout=20)
    runtime = json.loads(setup.stdout)
    assert runtime['openpyxl'] == '3.1.5', runtime
    initial_digest = digest(source)
    # Dependency isolation (venv) and the OS boundary are deliberately separate.
    # This fixed experiment checks write boundaries only. Startup needs macOS
    # runtime reads; a production backend must define its own read/network policy.
    profile = '\n'.join(['(version 1)', '(deny default)', '(allow process*)',
                         '(allow sysctl-read)', '(allow mach-lookup)',
                         '(allow file-read*)',
                         '(allow file-write* (subpath ' + json.dumps(str(outputs), ensure_ascii=False) + ')',
                         ' (subpath ' + json.dumps(str(temporary), ensure_ascii=False) + ') (literal "/dev/null"))'])
    (directory / 'sandbox.sb').write_text(profile)
    prefix = ['/usr/bin/sandbox-exec', '-p', profile, str(python)]
    started = time.monotonic()
    completed = subprocess.run(prefix + ['-c', '''
import json, sys
from pathlib import Path
from openpyxl import load_workbook
src,out,outside,link=map(Path,sys.argv[1:])
w=load_workbook(src, data_only=False); s=w['对账明细']
assert s['A2'].value=='00123' and s['B2'].value==30
assert s['C2'].value=='=B2*2' and s['A1'].font.bold
s['B2']=35; w.save(out)
r=load_workbook(out, data_only=False)['对账明细']
assert r['B2'].value==35 and r['C2'].value=='=B2*2'
assert r['A2'].value=='00123' and r['A1'].font.bold
checks={}
for name,path in [('inputReadOnly',src),('outsideOutputDenied',outside),('symlinkEscapeDenied',link)]:
 try:
  path.write_text('must not write', encoding='utf-8')
 except PermissionError:
  checks[name]=True
 else:
  raise AssertionError(name+' boundary failed')
checks.update({'nativeWorkbookReopened':True,'chinesePath':True,'formulaPreserved':True,
 'formulaCalculated':False,'renderingVerified':False})
print(json.dumps(checks))
''', str(source), str(target), str(sentinel), str(outputs / 'escape-link')],
        cwd=work, env=env, check=True, capture_output=True, text=True, timeout=20)
    facts = json.loads(completed.stdout)
    checked = subprocess.run(prefix + [str(checker), str(target), '--contains', '对账明细',
                                       '--contains', '00123', '--count', '1'],
                             cwd=work, env=env, check=True, capture_output=True, text=True, timeout=20)
    native_check = json.loads(checked.stdout)
    child = subprocess.Popen(prefix + ['-u', '-c', 'import time; print("started",flush=True); time.sleep(30)'],
                             cwd=work, env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                             text=True, start_new_session=True)
    try:
        import selectors
        selector = selectors.DefaultSelector()
        selector.register(child.stdout, selectors.EVENT_READ)
        assert selector.select(timeout=10), 'cancel probe did not start'
        assert child.stdout.readline().strip() == 'started'
        selector.close()
        cancel_start = time.monotonic()
        os.killpg(child.pid, signal.SIGTERM)
        child.communicate(timeout=3)
        facts['cancelStoppedProcess'] = child.returncode != 0
        facts['cancelMs'] = round((time.monotonic()-cancel_start)*1000)
        assert facts['cancelStoppedProcess']
    finally:
        if child.poll() is None:
            os.killpg(child.pid, signal.SIGKILL)
            child.communicate(timeout=3)
    assert digest(source) == initial_digest
    assert sentinel.read_text() == 'preserved'
    report = {'scope': 'MET164 PR1 feasibility only; not local Office product readiness',
              'platform': platform.platform(), 'runtime': runtime, 'facts': facts,
              'guideSha256': digest(guide), 'checkerSha256': digest(checker),
              'upstreamCheck': native_check, 'inputSha256': initial_digest,
              'outputSha256': digest(target), 'outputBytes': target.stat().st_size,
              'elapsedMs': round((time.monotonic()-started)*1000),
              'notCovered': ['macOS arm64', 'Bridge runtime installation/distribution',
                             'binary input/publication protocol', 'formula calculation',
                             'visual rendering', 'Word/PowerPoint', 'network/child-resource policy']}
    (directory / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(report, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except subprocess.CalledProcessError as error:
        print(json.dumps({'returncode':error.returncode, 'stderr':error.stderr}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)
