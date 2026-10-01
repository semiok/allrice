"""Fixed image-build checks/provenance. Never a tenant installer or execution path."""
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import platform
import shutil
import subprocess
import sys
import tempfile
import zipfile

BUILD = Path("/opt/allrice-build")
FONT = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
OFFICE = Path("/opt/dsh-office/scripts/check_office.py")
PNG = Path("/opt/allrice-python/check_png.py")
FONT_SHA = "b76b0433203017ca80401b2ee0dd69350349871c4b19d504c34dbdd80541690a"
OFFICE_SHA = "d94afa67593a284751e0f2dc000877e836f885d39954a882033cf22a13278f66"
PNG_SHA = "7109e40138f9089b3f260e66bc2b91fd7f1a90c57f3c43cbc2debf184c6a995f"


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def apt():
    lock = json.loads((BUILD / "apt.lock.json").read_text())
    assert subprocess.check_output(["dpkg", "--print-architecture"], text=True).strip() == lock["architecture"]
    for name, version in lock["basePackages"].items():
        actual = subprocess.check_output(["dpkg-query", "-W", "-f=${Version}", name], text=True)
        assert actual == version, (name, actual, version)
    assert len(lock["packages"]) == 1
    for package in lock["packages"]:
        raw = (BUILD / "debs" / package["fileName"]).read_bytes()
        assert len(raw) == package["sizeBytes"] and sha(raw) == package["sha256"]


def runtime(save=True):
    assert platform.python_version() == "3.11.13"
    assert sha(FONT.read_bytes()) == FONT_SHA
    assert sha(OFFICE.read_bytes()) == OFFICE_SHA
    assert sha(PNG.read_bytes()) == PNG_SHA
    architecture = {"aarch64": "arm64", "x86_64": "amd64"}.get(platform.machine())
    assert architecture is not None, ("unsupported architecture", platform.machine())
    apt_lock = json.loads(Path("/opt/allrice/apt.lock.json").read_text())
    assert architecture == apt_lock["architecture"]
    requirements = Path("/opt/allrice/requirements.lock").read_bytes()
    expected = {}
    for line in requirements.decode().splitlines():
        if line and not line.startswith(("#", " ")):
            name, version = line.rstrip(" \\").split("==")
            expected[name] = version
    packages = {name: importlib.metadata.version(name) for name in expected}
    assert packages == expected and len(packages) == 21
    if not save:
        private_cache = Path("/tmp/probe-mplconfig")
        shutil.copytree("/opt/python/mplconfig", private_cache)
        private_cache.chmod(0o700)
        for file in private_cache.iterdir():
            file.chmod(0o600)
        os.environ["MPLCONFIGDIR"] = str(private_cache)
    import matplotlib
    from matplotlib import font_manager
    import matplotlib.pyplot as plt
    assert matplotlib.get_backend().lower() == "agg"
    assert Path(font_manager.findfont("Noto Sans CJK JP", fallback_to_default=False)) == FONT
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pptx import Presentation
    from pptx.util import Inches
    with tempfile.TemporaryDirectory() as temp:
        directory = Path(temp)
        figure, axes = plt.subplots(figsize=(8, 4))
        axes.bar(["A001 办公费用", "A002 较长中文分类名称"], [10, -25])
        axes.set_title("中文图表与负数验收")
        axes.set_ylabel("金额（元）")
        chart = directory / "chart.png"
        figure.savefig(chart, format="png")
        plt.close(figure)
        raw = chart.read_bytes()
        checker = [sys.executable, "-I", str(PNG)]
        png_result = subprocess.run(checker, input=raw, capture_output=True, check=True, timeout=15)
        png_report = json.loads(png_result.stdout)
        assert png_report["checksum"] == "sha256:" + sha(raw)
        assert subprocess.run(checker, input=raw[:40], capture_output=True, timeout=15).returncode != 0
        document = Document()
        document.add_paragraph("中文费用报告：A001 10，A002 -25，合计 -15。")
        document.add_picture(str(chart))
        docx = directory / "proof.docx"
        document.save(docx)
        reopened = Document(docx)
        assert len(reopened.inline_shapes) == 1 and "合计 -15" in reopened.paragraphs[0].text
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["编号", "金额"])
        sheet.append(["A001", 10])
        sheet.append(["A002", -25])
        sheet.append(["合计", "=SUM(B2:B3)"])
        xlsx = directory / "proof.xlsx"
        workbook.save(xlsx)
        reopened_sheet = load_workbook(xlsx).active
        assert reopened_sheet["A2"].value == "A001" and reopened_sheet["B4"].value == "=SUM(B2:B3)"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "中文图表汇报"
        slide.shapes.add_picture(str(chart), Inches(1), Inches(1.5), width=Inches(8))
        pptx = directory / "proof.pptx"
        presentation.save(pptx)
        assert len(Presentation(pptx).slides) == 1
        for path in [docx, xlsx, pptx]:
            subprocess.run([sys.executable, "-I", str(OFFICE), str(path)], check=True, capture_output=True, timeout=15)
        for path in [docx, pptx]:
            with zipfile.ZipFile(path) as archive:
                media = [archive.read(name) for name in archive.namelist() if "/media/" in name]
                assert raw in media
    licenses = Path("/opt/allrice/licenses")
    font_license = Path("/usr/share/doc/fonts-noto-cjk/copyright")
    if save:
        shutil.copyfile(font_license, licenses / "Noto-CJK-copyright")
    source_licenses = json.loads((licenses / "source.lock.json").read_text())
    assert {item["name"] for item in source_licenses["packages"]} == {"et-xmlfile", "openpyxl"}
    package_licenses = {}
    for name in packages:
        dist = importlib.metadata.distribution(name)
        count = 0
        for file in dist.files or []:
            if not any(word in file.name.lower() for word in ["license", "copying", "notice", "copyright"]):
                continue
            source = dist.locate_file(file)
            if source.is_file():
                target = licenses / name / str(file)
                if save:
                    target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copyfile(source, target)
                else:
                    assert source.read_bytes() == target.read_bytes()
                count += 1
        origin = "wheel"
        if count == 0:
            source = next((item for item in source_licenses["packages"] if item["name"] == name), None)
            assert source is not None and source["version"] == packages[name], (name, "package license missing")
            raw = (licenses / source["fileName"]).read_bytes()
            assert len(raw) == source["sizeBytes"] and sha(raw) == source["sha256"]
            origin = "same_version_source_archive"
            count = 1
        package_licenses[name] = {"origin": origin, "files": count}
    proof = {
        "contractVersion": 1, "profileVersion": 1, "architecture": architecture,
        "pythonVersion": platform.python_version(),
        "packages": packages, "packagesChecksum": "sha256:" + sha(requirements),
        "aptChecksum": "sha256:" + sha(Path("/opt/allrice/apt.lock.json").read_bytes()),
        "font": {"fileName": FONT.name, "sha256": "sha256:" + FONT_SHA},
        "officeChecker": {"upstream": "@deepseek-ai/dsh-skill-office@0.1.7-alpha.2", "sha256": "sha256:" + OFFICE_SHA},
        "pngChecker": {"sha256": "sha256:" + PNG_SHA},
        "checks": {"cjkAggPng": True, "corruptPngRejected": True, "nativeOfficeReopen": True,
                   "dshOfficeCheck": True, "samePngEmbeddedInDocxPptx": True,
                   "excelFormulaPreserved": True},
        "packageLicenses": package_licenses,
        "licenseSourcesChecksum": "sha256:" + sha((licenses / "source.lock.json").read_bytes()),
    }
    if save:
        Path("/opt/allrice/runtime.json").write_text(json.dumps(proof, sort_keys=True, indent=2) + "\n")
    else:
        assert proof == json.loads(Path("/opt/allrice/runtime.json").read_text())
    print(json.dumps(proof, sort_keys=True))


if __name__ == "__main__":
    if sys.argv == [sys.argv[0], "apt"]:
        apt()
    elif sys.argv == [sys.argv[0], "runtime"]:
        runtime()
    elif sys.argv == [sys.argv[0], "probe"]:
        runtime(save=False)
    else:
        raise SystemExit("expected apt, runtime or probe")
